import sys
import pandas as pd
from openpyxl import load_workbook


def format_sheet(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col_cells in ws.columns:
        length = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max(length + 2, 10), 40)


base_stock = pd.read_csv("base_stock.csv")

try:
    wb = load_workbook("Sales_Pivot.xlsx")
    if "Base Stock" in wb.sheetnames:
        del wb["Base Stock"]
    ws = wb.create_sheet("Base Stock")
    ws.append(list(base_stock.columns))
    for row in base_stock.itertuples(index=False):
        ws.append(list(row))
    format_sheet(ws)
    wb.save("Sales_Pivot.xlsx")
    print("Appended Base Stock sheet to Sales_Pivot.xlsx")
except PermissionError:
    print("Sales_Pivot.xlsx is locked (open in Excel) - writing to Sales_Pivot_with_BaseStock.xlsx instead")
    sales = pd.read_csv("sales_pivot_step1.csv")
    with pd.ExcelWriter("Sales_Pivot_with_BaseStock.xlsx", engine="openpyxl") as writer:
        sales.to_excel(writer, sheet_name="Sales", index=False)
        format_sheet(writer.sheets["Sales"])
        base_stock.to_excel(writer, sheet_name="Base Stock", index=False)
        format_sheet(writer.sheets["Base Stock"])
    print("done")
